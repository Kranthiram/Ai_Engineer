import os
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq

load_dotenv()
my_api_key = os.getenv("GROQ_API_KEY")
if not my_api_key:
    raise ValueError("API Key not found")

client = Groq(api_key = my_api_key)
model = "openai/gpt-oss-120b"

#Part-1
job_description="""
Job Title: Java Full Stack Developer

Responsibilities
Develop and maintain web applications using Java and Spring Boot.
Design and develop RESTful APIs for communication between frontend and backend.
Build responsive frontend applications using React.js, JavaScript, HTML5 and CSS3.
Work with relational databases such as MySQL/PostgreSQL/Oracle and write SQL queries.
Implement application business logic using Java and Spring Boot.
Work with Spring Data JPA/Hibernate for database persistence.
Implement authentication and authorization using Spring Security/JWT.
Debug and troubleshoot application issues.
Write unit tests and participate in code reviews.
Use Git for source-code management.
Work with developers, testers, product teams and other stakeholders in an Agile/Scrum environment.
Participate in the complete SDLC: requirements → development → testing → deployment → support.

Required Skills
Core Java / Java 8+
OOP concepts
Collections, Exception Handling, Multithreading basics
Spring Framework
Spring Boot
Spring MVC
Spring Data JPA / Hibernate
REST APIs
Spring Security / JWT basics
SQL
MySQL/PostgreSQL
HTML, CSS, JavaScript
React.js or Angular
Git/GitHub
Maven/Gradle
Basic understanding of Microservices

Good to Have
Docker
Jenkins / GitHub Actions
AWS / Azure
Kafka
Redis
JUnit / Mockito
TypeScript
Kubernetes
 
Minimum Experience 0-2 years required and Education Requirements are any graducation degree in Computer science like B-tech,B.Com,IT,CS_AIML
"""

from pydantic import BaseModel, Field
class JobD(BaseModel):
    role:str
    required_skills:list[str]
    preferred_skills:list[str]
    minimum_experience: float | None
    education_requirements : list[str]
    responsibilities : list[str]

jobd_schema = JobD.model_json_schema()

system_prompt = f"""
You are an expert HR assistant.
Your job is to analyze job descriptions and extract structured information from them.

return ONLY valid JSON matching this schema:
{jobd_schema}

IMPORTANT
Do not return schema itself
Do not return fields like "properties", "title" or "type"
Fill the schema with actual information extracted from the job description.

If minimum experience is not mentioned, return null.
If information for a list is missing, return an empty list.
Do not invent information.
"""

user_prompt = f"""
Analyze the following job description:
{job_description}

"""

message_system = {
    "role":"system",
    "content": system_prompt
}

message_user = {
    "role": "user",
    "content":user_prompt
}

response_format={
    "type" : "json_object"
}

messages = [message_system, message_user]
response = client.chat.completions.create(model = model, messages = messages, response_format = response_format)
answer = response.choices[0].message.content
raw_json = answer

import json
job_data = json.loads(raw_json)
job = JobD(**job_data)

print("Minimum Experinece: ",job.minimum_experience)
print("Education Requirements: ",job.education_requirements)

#Part -2
class Experience(BaseModel):
    company : str | None = None
    role : str | None = None
    duration : str | None = None
    description : str | None = None
    skills_used : list[str] =[]

class Resume(BaseModel):
    name : str | None = None
    email : str | None = None
    phone : str | None = None

    total_experience_years : float | None = None

    skills : list[str] = []
    experiences : list[Experience] = []
    education : list[str] = []
    projects : list[str] = []
    certifications : list[str] = []

resume_schema = Resume.model_json_schema()

#Part -3
from pypdf import PdfReader
from docx import Document

def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text = text + page_text + "\n"
    return text

def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text

def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None

#part -4
def parse_resume(resume_text):
    system_prompt = f"""
    You are an expert resume parser

    Extract information from the resume based on its meaning,
    not only based on exact section headings.

    Different resumes may use different headings.

    For example:
    - Experience
    - Professional Experience
    - Work History
    - Employment
    - Internships

    These may all contain relevant experience.

    Skills may also appear in the skills section, work experience,
    internships or projects.

    Return ONLY valid JSON matching this schema:

    {resume_schema}

    Important rules:

    1. Do not invent information.
    2. If a value is not available, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
"""
    user_prompt = f"""
    Parse the following resume:
    {resume_text}
"""
    message_system = {
        "role": "system",
        "content":system_prompt
    }
    message_user={
        "role" : "user",
        "content" : user_prompt
    }
    messages=[message_system, message_user]
    response_format={
        "type": "json_object"
    }
    response=client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    raw_output = response.choices[0].message.content
    data  = json.loads(raw_output)
    resume = Resume(**data)
    return resume 

#Part -5
class MatchResult(BaseModel):
    score : float
    details : dict

def final_score(job,resume):
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
    You are an HR Recruiter
    Compare the candidate's resume with the job description

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}
    Retrun JSON matching this schema:

    {match_schema}

    Give me:
    1. Candidate name
    2. Matching skills
    3. Missing important skills
    4. Whether experience requirement is met
    5. Overall match percentage from 0 to 100
    6. A short final verdict

    Keep the response concise and easy to read.
    """
    message = {
        "role":"user",
        "content":prompt
    }

    messages = [message]
    response_format = {
        "type" : "json_object"
    }
    response = client.chat.completions.create(model = model, messages = messages, response_format = response_format)
    data = json.loads(response.choices[0].message.content)
    return MatchResult(**data)

#Part-6 FinalPart
import time
resume_folder = Path("resumes")
all_results = []

for file_path in resume_folder.iterdir():
    if file_path.suffix.lower() not in [".pdf",".docx"]:
        continue
    print("\nProcessing:", file_path)
    resume_text = read_resume(file_path)
    parsed_resume = parse_resume(resume_text) # LLM Call 1
    time.sleep(5)
    result = final_score(job,parsed_resume) # LLM Call 2
    time.sleep(5)
    print("Score: ",result.score)
    all_results.append({
        "name": parsed_resume.name,
        "score":result.score,
        "details":result.details
    })

all_results.sort(
    key = lambda candidate : candidate["score"],
    reverse = True
)

top_2 = all_results[:2]
worst_2 = all_results[-2:]

print("TOP 2 CANDIDATES")
for candidate in top_2:
    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )

    print(candidate["details"])

print("LOWEST 2 CANDIDATES")
for candidate in worst_2:

    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )
    print(candidate["details"])







